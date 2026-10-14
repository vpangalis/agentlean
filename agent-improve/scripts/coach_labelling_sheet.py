"""Founder ruling 2026-10-14, item 6 — the coach's replies with all A-fixes, on the 40 labelled turns, as a
labelling sheet (C: the regression set).

The 40 turns come from four walks of the same scripted Define run-through. Run 154211 is that walk again
with every A-fix (G-168, G-170 to G-175, G-176) — so for each labelled turn its reply is the turn of 154211
with the same Belt message at the same element (else the same message, else the same element and move).
Two labelled turns with the same input map to the same new reply; the sheet says so in `same as`.

Columns (founder): id, turn type, Belt message, coach reply, reply_good (yes/no), C1, C2, C3, C4
(pass / fail / n/a), note. Content decides; format faults go in the note.

    cd agent-improve; python -m scripts.coach_labelling_sheet [--run 20261013T154211]
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

PROJECT = Path(__file__).resolve().parents[1]
EVALS = PROJECT / "evals" / "define"
OUT = EVALS / "coach_labelling_sheet.xlsx"
BLOCKS = (("EXPLANATION", "explanation"), ("EXAMPLE - an illustration, not the Belt's data", "example"),
          ("QUESTION TO THE BELT", "prompt"))


def _norm(text: Any) -> str:
    return " ".join(str(text or "").split())


def shown(reply: dict[str, Any]) -> str:
    """The reply as the Belt sees it: the message, then the blocks, labelled."""
    parts = [str(reply.get("answer") or "").strip()]
    parts += [f"[{label}] {str(reply.get(key) or '').strip()}" for label, key in BLOCKS
              if str(reply.get(key) or "").strip()]
    return "\n\n".join(p for p in parts if p)


def match(turn: dict[str, Any], new: list[dict[str, Any]]) -> dict[str, Any] | None:
    field = (turn.get("script_step") or {}).get("field")
    belt = _norm(turn["belt"])

    def field_of(x: dict[str, Any]) -> Any:
        return x.get("field") or (x.get("reply") or {}).get("move_field")
    for pick in (lambda x: _norm(x.get("belt")) == belt and field_of(x) == field,
                 lambda x: _norm(x.get("belt")) == belt,
                 lambda x: field_of(x) == field and x.get("move") == turn["move"]):
        hits = [x for x in new if pick(x)]
        if hits:
            return hits[0]
    return None


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", default="20261013T154211")
    args = ap.parse_args()
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    turns = json.loads((EVALS / "judge_turns.json").read_text(encoding="utf-8"))["turns"]
    record = json.loads((PROJECT / "docs" / "runthrough" / f"define_runthrough_{args.run}.json").read_text(encoding="utf-8"))
    new = [x for x in record if isinstance(x, dict) and x.get("kind") == "turn"]
    kinds = {}
    old = EVALS / "judge_labelling_sheet.xlsx"
    if old.exists():
        ws = load_workbook(old, read_only=True)["Turns"]
        rows = list(ws.iter_rows(values_only=True))
        kinds = {str(r[0]): r[1] for r in rows[1:] if r and r[0]}

    wb = Workbook()
    about = wb.active
    about.title = "Why this file"
    for line in (f"Coach replies with all A-fixes — run {args.run}, the same scripted walk as the 40 labelled turns.",
                 "",
                 "reply_good: yes / no — content decides; format faults go in the note.",
                 "C1 Teaches before asking · C2 Challenges the right gap · C3 Faithful to the Belt · "
                 "C4 One clear next step — pass / fail / n/a (founder ruling 2026-10-14, grader v2).",
                 "'same as' names the first row with the same input: its reply is the same reply."):
        about.append([line])
    ws = wb.create_sheet("Turns")
    head = ["id", "turn type", "Belt message", "coach reply", "reply_good (yes/no)",
            "C1", "C2", "C3", "C4", "note (optional)", "new turn", "same as"]
    ws.append(head)
    first: dict[int, str] = {}
    unmatched = []
    for t in turns:
        x = match(t, new)
        if x is None:
            unmatched.append(t["id"])
            ws.append([t["id"], kinds.get(t["id"]), t["belt"], "(no matching turn in the new run)"])
            continue
        n = int(x["n"])
        same = first.get(n, "")
        first.setdefault(n, t["id"])
        move = str(x.get("move") or "")
        kind = move if move == t["move"] else f"{move} (was {t['move']})"
        ws.append([t["id"], kind, t["belt"], shown(x.get("reply") or {}), None,
                   None, None, None, None, None, f"{args.run}#{n}", same])
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    verdict = DataValidation(type="list", formula1='"pass,fail,n/a"', allow_blank=True)
    ws.add_data_validation(yes_no)
    ws.add_data_validation(verdict)
    last = ws.max_row
    yes_no.add(f"E2:E{last}")
    verdict.add(f"F2:I{last}")
    for col, width in zip("ABCDEFGHIJKL", (22, 22, 50, 90, 14, 9, 9, 9, 9, 40, 22, 22)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"
    wb.save(OUT)
    print(f"{len(turns)} rows, {len(set(first))} distinct new replies, unmatched {unmatched}; {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

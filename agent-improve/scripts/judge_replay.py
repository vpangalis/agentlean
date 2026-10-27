"""Founder ruling 2026-10-12 (quality stop, item 2): replay the two per-turn judges OFFLINE on the
evaluation set, and score them against the founder's labels.

For each turn in `evals/define/judge_turns.json` the coaching grader (the coherence judge is retired,
ADR-0083: its recorded verdicts are reported, never replayed)
are rebuilt for that turn's move and script step and asked about the reply they judged (a re-asked
turn's first reply) — the same calls the executor makes, nothing else: no graph, no run. The grader
is replayed on EVERY turn, coherence-rejected ones included, which is what DEF-015's option A will do.

Scored against `evals/define/judge_labels.json` ({turn id: {"reply_good": true|false, "rejection_right":
true|false|null, "note": "..."}}), per judge, for the verdicts the runs RECORDED and the verdicts REPLAYED now:
  false-reject rate = rejected replies labelled good / replies labelled good
  false-accept rate = accepted replies labelled not good / replies labelled not good
A turn with no label is replayed and reported, not scored. Calls, tokens, cost and time are reported.

    cd agent-improve; python -m scripts.judge_replay [--only <turn id> ...]
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
import time
from pathlib import Path
from typing import Any

EVALS = Path(__file__).resolve().parents[1] / "evals" / "define"
TURNS = EVALS / "judge_turns.json"
LABELS = EVALS / "judge_labels.json"
#: The founder's Excel sheet (2026-10-13): read when it carries labels; its labels are written to LABELS.
SHEET = EVALS / "judge_labelling_sheet.xlsx"
#: 2026-10-13: 5 rows of SHEET showed the wrong reply (the re-asked turns); their labels come from here.
CORRECTIONS = EVALS / "judge_labelling_corrections.xlsx"


def labels_from_sheet(path: Path) -> dict[str, dict[str, Any]]:
    """The Turns sheet's labels: reply_good yes/no -> True/False, judge_right right/wrong -> True/False,
    the optional note; a row with nothing filled is left out."""
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)["Turns"]
    rows = list(ws.iter_rows(values_only=True))
    col = {str(h or "").split(" ")[0]: i for i, h in enumerate(rows[0])}
    out: dict[str, dict[str, Any]] = {}
    for r in rows[1:]:
        good = str(r[col["reply_good"]] or "").strip().lower()
        right = str(r[col["judge_right"]] or "").strip().lower()
        note = str(r[col["note"]] or "").strip() if "note" in col else ""
        if not (good or right or note):
            continue
        out[str(r[col["id"]])] = {"reply_good": {"yes": True, "no": False}.get(good),
                                  "rejection_right": {"right": True, "wrong": False}.get(right),
                                  "note": note}
    return out
#: G-168 (founder ruling 2026-10-13): the planner's judgment, a third judge, on answers with a known verdict.
JUDGMENTS = EVALS / "judgment_cases.json"


async def replay_judgment(case: dict[str, Any]) -> dict[str, Any]:
    """The planner's one judgment on `case` — `nodes_common._judge`, as the planner calls it on a first
    answer (no earlier answer, not a read-back), the case's confirmed values empty."""
    from backend.phases import nodes_common
    t0 = time.time()
    from typing import cast

    from backend.core.substate import PhaseState
    state = cast(PhaseState, {"artifacts": {}, "phase_context": case.get("phase_context") or ""})
    j = await nodes_common._judge("define", state, case["field"], "", case["belt"], "no")
    return {"id": case["id"], "verdict": j.verdict, "failed_criterion": j.failed_criterion, "reason": j.reason,
            "expected": case["expected"], "agrees": j.verdict == case["expected"], "seconds": round(time.time() - t0, 1)}


#: Above this false-reject rate the founder's ruling stops the work for options (fix, merge, remove).
FALSE_REJECT_STOP = 0.20


def grader_for(turn: dict[str, Any]) -> Any:
    """The grader as the executor builds it for the turn: its move and its element's purpose (B2)."""
    from backend.middleware.grader import DMAICGraderMiddleware
    from backend.phases.nodes_common import _element_purpose
    field = (turn.get("script_step") or {}).get("field")
    return DMAICGraderMiddleware("define", move=turn["move"], purpose=_element_purpose("define", field))


async def replay(turn: dict[str, Any], turns: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """The grader, rebuilt for the turn's move, on the reply it judged (the coherence judge is retired,
    ADR-0083: its recorded verdicts are reported as history, never replayed)."""
    from backend.middleware.grader import DMAICGraderMiddleware
    t0 = time.time()
    grader = grader_for(turn)
    g = await grader._grade(turn["belt"], turn["reply_judged"])
    return {"id": turn["id"],
            "grader": {"passed": g.passed, "failed": [{"criterion": f.criterion, "reason": f.feedback} for f in g.failed],
                       "seconds": round(time.time() - t0, 1)}}


def rates(verdicts: dict[str, bool], labels: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """`verdicts`: turn id -> accepted. Scored against `reply_good` where it is labelled."""
    good = [i for i, l in labels.items() if i in verdicts and l.get("reply_good") is True]
    bad = [i for i, l in labels.items() if i in verdicts and l.get("reply_good") is False]
    fr = [i for i in good if not verdicts[i]]
    fa = [i for i in bad if verdicts[i]]
    return {"labelled": len(good) + len(bad),
            "false_reject_rate": round(len(fr) / len(good), 2) if good else None, "false_rejects": fr,
            "false_accept_rate": round(len(fa) / len(bad), 2) if bad else None, "false_accepts": fa}


def report(turns: list[dict[str, Any]], replays: list[dict[str, Any]], labels: dict[str, dict[str, Any]],
           cost: dict[str, Any], seconds: float) -> dict[str, Any]:
    by = {r["id"]: r for r in replays}
    recorded_coh = {t["id"]: bool(t["coherence"]["coherent"]) for t in turns}
    recorded_grd = {t["id"]: t["grader"]["status"] == "pass" for t in turns if t.get("grader")}
    replay_grd = {i: r["grader"]["passed"] for i, r in by.items()}
    rejections = [t["id"] for t in turns if not t["coherence"]["coherent"]]
    judged = [i for i in rejections if (labels.get(i) or {}).get("rejection_right") is not None]
    wrong = [i for i in judged if labels[i]["rejection_right"] is False]
    # Founder 2026-10-13: right to reject, for the wrong reason — a note on a rejection labelled right.
    rejected_any = [t["id"] for t in turns if not t["coherence"]["coherent"]
                    or (t.get("grader") or {}).get("status") == "failed"]
    wrong_reason = [i for i in rejected_any if (labels.get(i) or {}).get("rejection_right") is True
                    and str((labels.get(i) or {}).get("note") or "").strip()]
    out: dict[str, Any] = {
        "turns": len(turns), "replayed": len(replays),
        "coherence": {"recorded": rates(recorded_coh, labels), "retired": "ADR-0083"},
        "grader": {"recorded": rates(recorded_grd, labels), "replayed": rates(replay_grd, labels),
                   "replay_agrees_with_recorded": sum(replay_grd[i] == recorded_grd[i]
                                                      for i in replay_grd if i in recorded_grd)},
        "rejections_labelled_wrong": {"of": len(judged), "wrong": wrong,
                                      "share": round(len(wrong) / len(judged), 2) if judged else None},
        "right_outcome_wrong_reason": wrong_reason,
        "cost": cost, "seconds": round(seconds, 1),
        "per_turn_seconds": round(seconds / max(1, len(replays)), 1),
    }
    stop = [j for j in ("grader",)
            if (out[j]["replayed"]["false_reject_rate"] or 0) > FALSE_REJECT_STOP
            or (out[j]["recorded"]["false_reject_rate"] or 0) > FALSE_REJECT_STOP]
    out["stop"] = {"limit": FALSE_REJECT_STOP, "judges_over": stop}
    return out


#: ADR-0081's thresholds for leaving log-only: agreement with the founder's labels, and the false fails among the grader's.
AGREEMENT_MIN, FALSE_FAILS_MAX = 0.90, 0.20


def join_labels(labels: list[dict[str, str]], verdicts: list[dict[str, str]]) -> list[dict[str, str]]:
    """C10 (the founder's blind sheet, 2026-10-27): GRADER_LABELS.csv (the founder's pass / fail / n/a per criterion)
    joined to GRADER_VERDICTS.csv (the grader's verdict per criterion) by row id — one entry per row and criterion the
    founder labelled pass or fail and the grader judged pass or fail. A fail with no quote counts as the pass it was
    recorded as; "unknown", n/a and blanks are left out (an unknown is never a fail)."""
    by_id = {v["row_id"]: v for v in verdicts}
    out = []
    for row in labels:
        graded = by_id.get(row["row_id"]) or {}
        for key, label in row.items():
            if not key.endswith("_founder"):
                continue
            criterion = key[:-len("_founder")]
            verdict = str(graded.get(f"{criterion}_verdict") or "")
            verdict = "pass" if verdict.startswith("not counted") else verdict
            label = str(label or "").strip().lower()
            if label in ("pass", "fail") and verdict in ("pass", "fail"):
                out.append({"row": row["row_id"], "criterion": criterion, "grader_verdict": verdict,
                            "founder_label": label})
    return out


def label_agreement(rows: list[dict[str, str]]) -> dict[str, Any]:
    """C10 (founder ruling 2026-10-27): the grader's recorded verdicts against the founder's labels (`join_labels`),
    per criterion and overall — agreement on every labelled entry, on the grader's fails and on its passes, and the
    false fails (failed by the grader, labelled pass) — against ADR-0081's thresholds."""
    def score(group: list[dict[str, str]]) -> dict[str, Any]:
        fails = [r for r in group if r["grader_verdict"] == "fail"]
        passes = [r for r in group if r["grader_verdict"] == "pass"]
        agree = [r for r in group if r["grader_verdict"] == r["founder_label"]]
        false_fails = [r for r in fails if r["founder_label"] == "pass"]
        return {"labelled": len(group), "agree": len(agree),
                "agreement": round(len(agree) / len(group), 3) if group else None,
                "fails": len(fails), "fails_agreed": len(fails) - len(false_fails),
                "false_fail_rate": round(len(false_fails) / len(fails), 3) if fails else None,
                "passes": len(passes), "passes_agreed": sum(r["founder_label"] == "pass" for r in passes),
                "disagreements": [r["row"] for r in group if r["grader_verdict"] != r["founder_label"]]}
    per = {c: score([r for r in rows if r["criterion"] == c]) for c in sorted({r["criterion"] for r in rows})}
    overall = score(rows)
    met = bool(rows) and (overall["agreement"] or 0) >= AGREEMENT_MIN \
        and (overall["false_fail_rate"] or 0) < FALSE_FAILS_MAX
    return {"overall": overall, "per_criterion": per, "thresholds": {"agreement": AGREEMENT_MIN,
            "false_fails_below": FALSE_FAILS_MAX}, "met": met}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    ap = argparse.ArgumentParser()
    ap.add_argument("--labels", default="", help="GRADER_LABELS.csv with the founder's labels: agreement per criterion")
    ap.add_argument("--verdicts", default="", help="GRADER_VERDICTS.csv, joined to --labels by row id")
    ap.add_argument("--only", nargs="*", default=None)
    ap.add_argument("--judgments-only", action="store_true", help="replay only the planner-judgment cases")
    args = ap.parse_args()
    if args.labels:
        import csv
        with open(args.labels, encoding="utf-8-sig", newline="") as lf, \
                open(args.verdicts, encoding="utf-8-sig", newline="") as vf:
            joined = join_labels(list(csv.DictReader(lf)), list(csv.DictReader(vf)))
        print(json.dumps(label_agreement(joined), indent=1, ensure_ascii=False))
        return 0
    from scripts.define_runthrough import _cost
    from scripts.move_proof_661 import CALLS, _count_model_calls, token_totals
    from scripts.coaching_proof_656 import _tracing_off
    _tracing_off()
    _count_model_calls(200)
    judgment_cases = json.loads(JUDGMENTS.read_text(encoding="utf-8"))["cases"] if JUDGMENTS.exists() else []
    if args.judgments_only:
        started = time.time()

        async def judged() -> list[dict[str, Any]]:
            return [await replay_judgment(c) for c in judgment_cases]
        results = asyncio.run(judged())
        print(json.dumps({"judgments": results, "agree": sum(r["agrees"] for r in results), "of": len(results),
                          "cost": _cost(token_totals(CALLS)), "seconds": round(time.time() - started, 1)},
                         indent=1, ensure_ascii=False))
        return 0
    turns = json.loads(TURNS.read_text(encoding="utf-8"))["turns"]
    every = list(turns)
    if args.only:
        turns = [t for t in turns if t["id"] in set(args.only)]
    sheet = labels_from_sheet(SHEET) if SHEET.exists() else {}
    if CORRECTIONS.exists():
        sheet.update(labels_from_sheet(CORRECTIONS))
    if sheet:
        LABELS.write_text(json.dumps(sheet, indent=1, ensure_ascii=False) + chr(10), encoding="utf-8")
    labels = json.loads(LABELS.read_text(encoding="utf-8")) if LABELS.exists() else {}
    started = time.time()

    async def all_turns() -> list[dict[str, Any]]:
        return [await replay(t, every) for t in turns]
    replays = asyncio.run(all_turns())

    async def judged_all() -> list[dict[str, Any]]:
        return [await replay_judgment(c) for c in judgment_cases]
    judgments = asyncio.run(judged_all()) if not args.only else []
    seconds = time.time() - started
    out = report(turns, replays, labels, _cost(token_totals(CALLS)), seconds)
    out["planner_judgment"] = {"cases": len(judgments), "agree_with_expected": sum(j["agrees"] for j in judgments),
                               "disagree": [j["id"] for j in judgments if not j["agrees"]]}
    stamp = time.strftime("%Y%m%dT%H%M%S", time.gmtime())
    path = EVALS / f"judge_replay_{stamp}.json"
    path.write_text(json.dumps({"summary": out, "replays": replays, "judgments": judgments}, indent=1,
                               ensure_ascii=False) + "\n",
                    encoding="utf-8")
    print(json.dumps(out, indent=1, ensure_ascii=False))
    print(f"record: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
